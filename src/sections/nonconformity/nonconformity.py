from docx.shared import Pt, RGBColor, Inches
from docx.document import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from typing import Any, Dict, List
import pandas as pd
import re

from utils import (
    adicionar_paragrafo_justificado,
    adicionar_titulo_secao,
    adicionar_paragrafo_info_compacta,
    aplicar_estilo_corpo,
    aplicar_estilo_titulo,
    encontrar_dados_na_base,
    LARGURA_PADRAO_IN
)

def _limpar_inicio_texto(texto: str) -> str:
    if not texto: return ""
    s = str(texto).strip()
    padrao = r'^(?:[A-Z]{3}[\s\-]*)?\d+(?:[._]\d+)?\s*[-:–)]*\s*'
    texto_limpo = re.sub(padrao, '', s)
    if not texto_limpo and s: return s
    if texto_limpo and texto_limpo[0].islower():
        return texto_limpo[0].upper() + texto_limpo[1:]
    return texto_limpo

def _inserir_texto_nc(doc: Document, nc_titulo_identificador: str, descricao_bruta: str):
    if not descricao_bruta: descricao_bruta = "Descrição indisponível."
    linhas = [l.strip() for l in descricao_bruta.split("\n") if l.strip()]
    if not linhas: linhas = ["Descrição indisponível."]

    paragrafo_nc = doc.add_paragraph()
    run_titulo = paragrafo_nc.add_run(f"Não Conformidade {nc_titulo_identificador}")
    aplicar_estilo_corpo(run_titulo, negrito=True)
    run_titulo.underline = True

    run_traco = paragrafo_nc.add_run(" - ")
    aplicar_estilo_corpo(run_traco)

    texto_principal = _limpar_inicio_texto(linhas[0])
    run_desc = paragrafo_nc.add_run(texto_principal)
    aplicar_estilo_corpo(run_desc)

    paragrafo_nc.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY_LOW

    base_prefixo = ""
    match_base = re.match(r"([A-Z]{3}\s\d{4}_)", nc_titulo_identificador)
    if match_base: base_prefixo = match_base.group(1)

    if len(linhas) > 1:
        for sub_item_raw in linhas[1:]:
            p_sub = doc.add_paragraph()
            p_sub.paragraph_format.left_indent = Inches(0.5)
            p_sub.paragraph_format.space_after = Pt(2)
            p_sub.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY_LOW

            match_sub = re.match(r'^(?:[A-Z]{3}[\s\-]*)?(\d+[._]\d+)\s*[-:–]?\s*(.*)', sub_item_raw)

            if match_sub:
                numero_sub = match_sub.group(1).replace("_", ".") 
                texto_restante = match_sub.group(2)
                texto_final = _limpar_inicio_texto(texto_restante)
                id_completo = f"{base_prefixo}{numero_sub}" if base_prefixo else f"Item {numero_sub}"

                run_sub_id = p_sub.add_run(f"Não Conformidade {id_completo}")
                aplicar_estilo_corpo(run_sub_id, negrito=True)
                p_sub.add_run(" - ")
                run_txt = p_sub.add_run(texto_final)
                aplicar_estilo_corpo(run_txt)
            else:
                texto_limpo = _limpar_inicio_texto(sub_item_raw)
                run_txt = p_sub.add_run(texto_limpo)
                aplicar_estilo_corpo(run_txt)

def _inserir_linhas_info(doc: Document, info_socicam: str, constatacao_monit: str, analise_arpe: str):
    vals = {}
    for var_name, value in [("info_socicam", info_socicam), ("constatacao_monit", constatacao_monit), ("analise_arpe", analise_arpe)]:
        if str(value).strip().lower() in ["nan", "none", "", "nat"]:
             vals[var_name] = "N/A"
        else:
             vals[var_name] = str(value).strip()
             
    adicionar_paragrafo_info_compacta(doc, "Informação da SOCICAM: ", vals["info_socicam"])
    adicionar_paragrafo_info_compacta(doc, "Constatação: ", vals["constatacao_monit"])
    adicionar_paragrafo_info_compacta(doc, "Análise da ARPE: ", vals["analise_arpe"])

def _safe_split(texto: Any) -> List[str]:
    s = str(texto).strip()
    if not s or s.lower() == "nan": return []
    partes = re.split(r'[;\n]|:(?!\d)', s)
    return [x.strip() for x in partes if x.strip()]

def gerar_secao_nao_conformidades_constatadas(
    doc: Document, row: pd.Series, nao_conformidades_df: pd.DataFrame, FOTOS_DIR: str,
    processo_info: Dict[str, str], df_base_nc: pd.DataFrame, ano_user: str, proc_user: str, monit_user: str
):
    id_fiscalizacao = row["ID da Fiscalização"]
    processo_ctr = processo_info.get("Processo CTR Nº", "XX/XXXX")
    cartas_raw = str(processo_info.get("Carta SAP/PER/ARPE Nº", "")).strip()
    texto_cartas = ""
    
    if cartas_raw and cartas_raw.lower() != "nan":
        partes_cartas = [p.strip() for p in cartas_raw.split(";") if p.strip()]
        if partes_cartas:
            cartas_formatadas = [f"Carta SAP/PER/ARPE N° {p}" for p in partes_cartas]
            juncao_cartas = " e ".join(cartas_formatadas)
            texto_cartas = f"constante da {juncao_cartas}, "

    adicionar_titulo_secao(doc, "3. RESULTADO DAS VISTORIAS DAS NÃO CONFORMIDADES PENDENTES")
    adicionar_paragrafo_justificado(
        doc, ("Estão registrados para cada Terminal Rodoviário os resultados da verificação pela Arpe das ações "
              f"desenvolvidas pela SOCICAM, {texto_cartas}para solucionar as Não Conformidades ainda pendentes "
              f"apresentadas no Relatório de Fiscalização Técnico-Operacional ARPE/CTR nº {processo_ctr}.")
    )

    grupo = nao_conformidades_df[nao_conformidades_df["ID da Fiscalização"] == id_fiscalizacao].copy()
    grupos_ordenados = sorted(grupo.groupby("Terminal"), key=lambda x: str(x[0]))
    
    num_terminal = 1
    for terminal, dados_terminal in grupos_ordenados:
        p = doc.add_paragraph()
        run = p.add_run(f"3.{num_terminal} - {str(terminal).upper()}")
        aplicar_estilo_titulo(run)
        p.paragraph_format.space_before = Pt(12)

        ncs_processadas = set()
        for _, nc in dados_terminal.iterrows():
 
            raw_legenda = str(nc.get("Legenda da Foto", "")).strip()
            raw_const_display = str(nc.get("Constatação", "")).strip()
            raw_socicam = str(nc.get("Informação SOCICAM", "")).strip()
            raw_arpe = str(nc.get("Análise da Arpe", "")).strip()

            list_legenda = _safe_split(raw_legenda)

            list_const_display = _safe_split(raw_const_display)
            list_socicam = _safe_split(raw_socicam)
            list_arpe = _safe_split(raw_arpe)

            # Define chaves de busca
            if list_legenda:
                search_keys = list_legenda
            else:
                search_keys = list_const_display 

            # Índice separado para controlar os textos de exibição
            idx_exibicao = 0 
            
            for txt_busca in search_keys:
                if not txt_busca: continue

                dados = encontrar_dados_na_base(txt_busca, df_base_nc, ano_user, proc_user, monit_user, terminal_user=str(terminal))

                # Se não encontrou, PULA e NÃO consome o índice de exibição
                if "NAO_ENCONTRADO" in str(dados["id"]): 
                    continue

                if dados["id"] in ncs_processadas: continue
                ncs_processadas.add(dados["id"])
                
                txt_exibir_const = list_const_display[idx_exibicao] if idx_exibicao < len(list_const_display) else "N/A"
                txt_exibir_socicam = list_socicam[idx_exibicao] if idx_exibicao < len(list_socicam) else "N/A"
                txt_exibir_arpe = list_arpe[idx_exibicao] if idx_exibicao < len(list_arpe) else "N/A"
                
                # Incrementa o índice apenas porque USAMOS o texto
                idx_exibicao += 1

                _inserir_texto_nc(doc, dados["id"], dados["descricao"])
                _inserir_linhas_info(doc, txt_exibir_socicam, txt_exibir_const, txt_exibir_arpe)
                
                doc.add_paragraph().paragraph_format.space_after = Pt(12)

        num_terminal += 1