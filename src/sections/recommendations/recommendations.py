from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from utils import adicionar_titulo_secao, adicionar_paragrafo_justificado 

def gerar_secao_recomendacoes(doc: Document, row):
    """
    Gera a seção '6. RECOMENDAÇÕES', formatando os itens como uma lista de bullet points.
    """
    
    adicionar_titulo_secao(doc, "6. RECOMENDAÇÕES")
    doc.add_paragraph() 

    texto_intro = (
        "Considerando as disposições do Contrato de Concessão, em especial, o Anexo III – Regulamento Interno dos Terminais "
        "Rodoviários, aprovado pela Resolução Arpe nº 46, de 07 de abril de 2008 (Antiga nº 06/2008), bem como a legislação "
        "aplicável, devem ser observadas pela SOCICAM as seguintes recomendações:"
    )
    
    adicionar_paragrafo_justificado(doc, texto_intro)

    recomendacoes = [
        
        "Garantir condições de segurança, higiene, acessibilidade e conforto aos usuários dos Terminais Rodoviários, "
        "sejam passageiros, público em geral, comerciantes neles estabelecidos, empresas de transportes e de seus empregados.",
      
        "Exigir a utilização de EPI adequados, inclusive por funcionários de empresas terceirizadas que prestem serviços nos Terminais Rodoviários.",
  
        "Providenciar a correta manutenção (evitar o vencimento) de extintores de incêndios nos Terminais Rodoviários.",

        "Instalar, sempre que necessário, aviso de sinalização de segurança, principalmente em pontos de risco de acidentes.",
        
        "Levantar a necessidade de manutenção das cobertas de todos os Terminais Rodoviários, em especial, placas cimentícias das "
        "testeiras e vigas, considerando que quatro das nove NC registradas neste Relatório apontam "
        "para essa questão. Assim, solicita-se laudos técnicos sobre a situação das cobertas, suas condições estruturais "
        "e funcionalidade de drenagem pluvial."
    ]

    for item in recomendacoes:
        par_item = doc.add_paragraph(item, style='List Bullet')

        par_item.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    doc.add_paragraph()