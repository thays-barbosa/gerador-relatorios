import os
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.shared import Inches, Pt 

from utils import (
    parear_responsaveis_e_matriculas, 
    formatar_data_capa, 
    adicionar_texto_centralizado,
    remover_espacamento_paragrafo, 
    adicionar_quebra_linha_controlada 
)


def gerar_capa(doc: Document, caminho_base_dir, info_fiscalizacao):
    """
    Gera a capa do relatório de forma dinâmica, extraindo dados da planilha,
    com controle de espaçamento para manter o layout compacto e centralizado em uma página,
    forçando o título superior para o topo e aplicando negrito conforme solicitado.
    """
  
    
    nomes_responsaveis = info_fiscalizacao.get('Pessoal Responsável', info_fiscalizacao.get('PESSOAL RESPONSÁVEL', ''))
    matriculas_responsaveis = info_fiscalizacao.get('Matrícula do Pessoal Responsável', '') 
    profissionais = parear_responsaveis_e_matriculas(nomes_responsaveis, matriculas_responsaveis)

    data_formatada = formatar_data_capa(str(info_fiscalizacao.get('Data', ''))) 
    num_processo = info_fiscalizacao.get('Nº Processo', 'N/A')
    num_sei = info_fiscalizacao.get('Nº SEI', 'N/A')
    
    relatorio_fixo = f"RELATÓRIO DE FISCALIZAÇÃO PROC ADM Nº {num_processo} - CTR"
    sei_fixo = f"SEI Nº {num_sei}"
    
    ALTURA_QUEBRA_PT = 15 
    
    QUEBRAS_POS_TITULO = 1
    QUEBRAS_POS_LOGO = 1
    QUEBRAS_PARA_ASSINATURAS = 1
    QUEBRAS_POS_DATA = 1

    
    if len(doc.paragraphs) > 0 and doc.paragraphs[0].text == '':
        par_titulo = doc.paragraphs[0]
    else:
        par_titulo = doc.add_paragraph() 
        
    remover_espacamento_paragrafo(par_titulo) 
    par_titulo.paragraph_format.line_spacing_rule = None
    par_titulo.paragraph_format.line_spacing = None
    
    par_titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par_titulo.text = ''
    
    run_titulo = par_titulo.add_run("RELATÓRIO DE FISCALIZAÇÃO")
    run_titulo.bold = True
    run_titulo.font.size = Pt(12) 

    for _ in range(QUEBRAS_POS_TITULO):
         adicionar_quebra_linha_controlada(doc, altura_pt=ALTURA_QUEBRA_PT)

    try:
        caminho_logo = os.path.join(caminho_base_dir, "assets/logo_arpe.jpg") 
        if not os.path.exists(caminho_logo):
            raise FileNotFoundError(f"Arquivo de logo não encontrado no caminho: {caminho_logo}")
            
        doc.add_picture(caminho_logo, width=Inches(6.0)) 
        logo_arpe = doc.paragraphs[-1]
        
        remover_espacamento_paragrafo(logo_arpe)
        logo_arpe.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
    except Exception as e:
        print(f"\n⚠️ ERRO GRAVE AO INSERIR LOGO NA CAPA: {e}") 
        doc.add_paragraph("--- FALHA AO CARREGAR LOGO ARPE (Verifique /assets/logo_arpe.jpg) ---").alignment = WD_ALIGN_PARAGRAPH.CENTER
        
 
    for _ in range(QUEBRAS_POS_LOGO): 
        adicionar_quebra_linha_controlada(doc, altura_pt=ALTURA_QUEBRA_PT) 
  
    par_fisca = doc.add_paragraph("FISCALIZAÇÃO NOS TERMINAIS RODOVIÁRIOS INTERMUNICIPAIS DE PASSAGEIROS")
    remover_espacamento_paragrafo(par_fisca)
    par_fisca.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par_fisca.runs[0].bold = True 
  
    par_prestador = doc.add_paragraph("PRESTADOR DE SERVIÇO: SOCICAM - ADMINISTRAÇÃO, PROJETOS E REPRESENTAÇÕES LTDA")
    remover_espacamento_paragrafo(par_prestador)
    par_prestador.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par_prestador.runs[0].bold = True 
   
    for _ in range(QUEBRAS_PARA_ASSINATURAS): 
        adicionar_quebra_linha_controlada(doc, altura_pt=ALTURA_QUEBRA_PT) 

    for profissional in profissionais:
  
        par_nome = doc.add_paragraph()
        remover_espacamento_paragrafo(par_nome) 
        par_nome.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par_nome.add_run(profissional['nome']).bold = True

        par_info = doc.add_paragraph()
        remover_espacamento_paragrafo(par_info) 
        par_info.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par_info.add_run(profissional['info_completa'])
  
        adicionar_quebra_linha_controlada(doc, altura_pt=5) 
   
    adicionar_quebra_linha_controlada(doc, altura_pt=ALTURA_QUEBRA_PT) 
  
    par_data = doc.add_paragraph()
    remover_espacamento_paragrafo(par_data) 
    par_data.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par_data.add_run(data_formatada).bold = True 
    
    for _ in range(QUEBRAS_POS_DATA): 
        adicionar_quebra_linha_controlada(doc, altura_pt=ALTURA_QUEBRA_PT) 

    par_relatorio = doc.add_paragraph()
    remover_espacamento_paragrafo(par_relatorio) 
    par_relatorio.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par_relatorio.add_run(relatorio_fixo).bold = True 
    
    par_sei = doc.add_paragraph()
    remover_espacamento_paragrafo(par_sei) 
    par_sei.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par_sei.add_run(sei_fixo).bold = True 
   
    doc.add_section(WD_SECTION.NEW_PAGE)