from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from utils import (
    adicionar_titulo_secao,
    adicionar_paragrafo_justificado,
   
)

def gerar_secao_determinacoes_finais(doc: Document, row):
    """
    Gera a seção '5. DETERMINAÇÕES GERAIS' do relatório, formatada com 
    subtítulos em negrito conforme o modelo.
    """

    adicionar_titulo_secao(doc, "5. DETERMINAÇÕES GERAIS")

    doc.add_paragraph() 
  
    texto1 = (
        "Considerando os dispositivos contratuais pertinentes e visando garantir a qualidade dos serviços prestados, "
        "determina-se que a SOCICAM tome as seguintes medidas através de um plano de ação:"
    )
   
    adicionar_paragrafo_justificado(doc, texto1)

    par2 = doc.add_paragraph()
    par2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    par2.add_run("Manutenção e Monitoramento").bold = True
 
    par2.add_run(
        ": adotar medidas para assegurar a manutenção, o monitoramento contínuo e o cumprimento do "
        "Programa de Manutenção dos Terminais Rodoviários, constante da proposta da SOICICAM nos subitens 9.1.1  "
        "Manutenção Preventiva; 9.1.2  manutenção Corretiva e 9.13  tabela de classificação de níveis de falha "
        "(tabela de tempos máximos para os níveis de atendimento)."
    )

    par3 = doc.add_paragraph()
    par3.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

  
    par3.add_run("Medidas imediatas").bold = True
    
   
    par3.add_run(
        " para resolutividade das "
    )

    par3.add_run("NC ").bold = True 

    par3.add_run("constatadas, nos prazos estabelecidos, conforme disposto no Quadro 1, na coluna denominada Determinações. ")

    doc.add_paragraph()