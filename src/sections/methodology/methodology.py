from docx import Document
from utils import adicionar_titulo_secao, adicionar_paragrafo_justificado
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches 

def gerar_secao_metodologia(doc: Document, row):
    """
    Adiciona a seção "3. METODOLOGIA" ao documento, formatada conforme o modelo
    (subtítulos em negrito e lista de bullet points).
    """
  
    adicionar_titulo_secao(doc, "3. METODOLOGIA")
 
    doc.add_paragraph() 

    texto_intro = (
        "A fiscalização direta e periódica realizada pela Coordenadoria de Transportes e Rodovias da Arpe está submetida a uma "
        "metodologia organizada em três etapas: Preparação e Planejamento, Execução da Fiscalização e Monitoramento e Avaliação."
    )
    
    adicionar_paragrafo_justificado(doc, texto_intro)

    par_prep = doc.add_paragraph()
   
    par_prep.add_run("Preparação e Planejamento").bold = True
    par_prep.add_run(
        " - compreende a organização e estruturação das atividades preliminares à execução da "
        "fiscalização, destacando-se a elaboração e o envio de avisos de fiscalização à Concessionária e demais atividades de "
        "suporte à fiscalização, bem como a análise de fiscalizações anteriores com a identificação de eventuais Não Conformidades pendentes. "
    )
    
    par_prep.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    par_exec = doc.add_paragraph()
    par_exec.add_run("Execução da Fiscalização").bold = True
    par_exec.add_run(
        " - a execução da fiscalização é pautada por um arcabouço de normas e diretrizes, " 
        "possibilitando que todas as etapas sejam desenvolvidas de maneira eficiente e em conformidade aos padrões "
        "estabelecidos, destacando-se:"
    )
    par_exec.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    itens_normas = [
        "Lei nº 13.254, de 21 de junho de 2007, alterada pela Lei nº 15.200, de 17 de dezembro de 2013, e regulamentada pelo Decreto nº 40.559, de 31 de março de 2014.",
        "Resoluções Arpe nº 46, de 07 de abril de 2008 (Antiga nº 06/2008), alterada pela Resolução ARPE nº 53, de 26 de janeiro de 2009 (Antiga 003/2009); e nº 083, de 30 de julho de 2013.",
        "Contrato de Concessão de Serviço Público Nº 1.041.080/08, de 19 de setembro de 2008 e aditivos, em especial, o Segundo Termo Aditivo ao Contrato de Concessão, de 29 de setembro de 2017.",
        "Normas Técnicas da ABNT."
    ]

    for item in itens_normas:
        par_item = doc.add_paragraph(item, style='List Bullet')

    doc.add_paragraph() 

    par_monitor = doc.add_paragraph()
    par_monitor.add_run("Monitoramento e Avaliação").bold = True
    par_monitor.add_run(
        " - Esta etapa é fundamental para garantir a eficácia das ações corretivas a serem "
        "executadas pela Concessionária para a melhoria contínua dos serviços prestados. Os principais instrumentos do " 
        "Monitoramento e Avaliação são: Termo de Notificação e respectivo Relatório de Fiscalização, Plano de Ação da "
        "Concessionária e Relatórios de Monitoramento e Avaliação Final. " 
    )
    par_monitor.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY