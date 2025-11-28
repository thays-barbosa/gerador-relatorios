from typing import List, Dict, Any
import pandas as pd
from docx.document import Document
from docx.shared import Pt, Inches
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml.shared import OxmlElement
import re

from utils import (
    adicionar_titulo_secao,
    aplicar_estilo_corpo,
    aplicar_estilo_paragrafo_compacto,
    encontrar_dados_na_base
)

def _aplicar_cor_fundo_celula(cell, cor_hex: str = "D9D9D9"):
    tc = cell._element
    tcPr = tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), cor_hex)
    tcPr.append(shading)

def _aplicar_estilo_resumo(run, negrito: bool = False):
    aplicar_estilo_corpo(run, negrito=negrito)

def _formatar_nome_terminal(nome_bruto: str) -> str:
    return str(nome_bruto).upper().replace("TERMINAL DE ", "").replace("TERMINAL DO ", "").strip()

def _safe_split_resumo(texto: Any) -> List[str]:
    """
    Divide por ';' ou Enter ou ':' (protegendo horários como 10:30).
    """
    s = str(texto).strip()
    if not s or s.lower() == 'nan': return []

    partes = re.split(r'[;\n]|:(?!\d)', s)
    
    return [x.strip() for x in partes if x.strip()]

def _limpar_redundancia_tabela(texto: str) -> str:
    """Remove 'TIP 01', 'CAR 05' etc. do início."""
    if not texto: return ""
    padrao = r'^([A-Z]{3}\s+)?\d+([._]\d+)?\s*[-:–]?\s*'
    limpo = re.sub(padrao, '', str(texto).strip())
    
    if not limpo and texto: return str(texto)

    if limpo and limpo[0].islower():
        return limpo[0].upper() + limpo[1:]
    return limpo

def gerar_secao_resumo_nao_conformidades(
    doc: Document,
    row: pd.Series,
    nao_conformidades_df: pd.DataFrame,
    processo_info: Dict[str, str],
    df_base_nc: pd.DataFrame,
    ano_user: str,
    proc_user: str,
    monit_user: str
):
    id_fisc = row["ID da Fiscalização"]
    processo_ctr = processo_info.get("Processo CTR Nº", "XX/XXXX")
    carta_bruta = str(processo_info.get("Carta SAP/PER/ARPE Nº", "XXX/XXXX"))
    carta_limpa = carta_bruta.replace(";", " ").strip()
    periodo_bruto = str(processo_info.get("Periodo de Vistoria da ARPE", "DATA INDEFINIDA"))
    periodo_vistoria = periodo_bruto.replace(";", " e ")

    nc_fisc = nao_conformidades_df[nao_conformidades_df["ID da Fiscalização"] == id_fisc].copy()
    if nc_fisc.empty:
        doc.add_paragraph("Nenhuma não conformidade registrada.")
        return

    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    adicionar_titulo_secao(doc, "4. RESUMO DA SITUAÇÃO DAS NÃO CONFORMIDADES MONITORADAS")
    
    tabela = doc.add_table(rows=1, cols=5)
    tabela.style = "Table Grid"
    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    headers = [
        "TERMINAL",
        f"NÃO CONFORMIDADE\nRELATÓRIO ARPE/CTR\n{processo_ctr}",
        f"INFORMAÇÃO SOCICAM\nCarta SAP/PER/ARPE\n{carta_limpa}",
        f"VISTORIA DA ARPE\n{periodo_vistoria}",
        "SITUAÇÃO"
    ]
    col_widths = [Inches(1.0), Inches(2.2), Inches(1.9), Inches(1.5), Inches(0.8)]
    
    for i, titulo in enumerate(headers):
        cell = tabela.rows[0].cells[i]
        cell.text = titulo
        _aplicar_cor_fundo_celula(cell)
        cell.width = col_widths[i]
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if cell.paragraphs[0].runs:
            _aplicar_estilo_resumo(cell.paragraphs[0].runs[0], negrito=True)

    current_row_index = 1
    grupos_terminais = sorted(nc_fisc.groupby("Terminal"), key=lambda x: str(x[0]))

    for terminal_bruto, grupo in grupos_terminais:
        primeira_celula_terminal = None
        nome_terminal = _formatar_nome_terminal(terminal_bruto)
        lista_dados = []
        ids_adicionados = set()

        for _, linha in grupo.iterrows():
            raw_const = str(linha.get("Legenda da Foto", "")).strip()
            if not raw_const or raw_const.lower() == "nan":
                 raw_const = str(linha.get("Constatação", "")).strip()

            constatacoes = _safe_split_resumo(raw_const)
            
            raw_info = str(linha.get("Informação SOCICAM carta", "")).strip()
            if not raw_info or raw_info.lower() == 'nan':
                raw_info = str(linha.get("Informação SOCICAM", "")).strip()
            
            infos = _safe_split_resumo(raw_info)
            
            max_len = max(len(constatacoes), len(infos))
            if len(constatacoes) < max_len: constatacoes.extend([""] * (max_len - len(constatacoes)))
            if len(infos) < max_len: infos.extend(["N/A"] * (max_len - len(infos)))

            for i in range(max_len):
                txt_busca = constatacoes[i]
                if not txt_busca: continue
                
                info_socicam_texto = infos[i]
                
                # AQUI ELE CHAMA A FUNÇÃO CORRIGIDA NO UTILS.PY
                dados_base = encontrar_dados_na_base(
                    txt_busca, df_base_nc, ano_user, proc_user, monit_user, terminal_user=str(terminal_bruto)
                )
                
                if dados_base["id"] in ids_adicionados and dados_base["id"] != "ID_NAO_ENCONTRADO":
                    continue
                
                if dados_base["id"] != "ID_NAO_ENCONTRADO":
                    ids_adicionados.add(dados_base["id"])
                
                lista_dados.append({
                    "id": dados_base["id"],
                    "desc": dados_base["descricao"],
                    "socicam": info_socicam_texto,
                    "data": dados_base["data_vistoria"],
                    "situacao": dados_base["situacao"]
                })

        lista_dados.sort(key=lambda x: x["id"])

        if not lista_dados: continue

        for idx, item in enumerate(lista_dados):
            row_cells = tabela.add_row().cells
            
            if idx == 0:
                primeira_celula_terminal = row_cells[0]
                row_cells[0].text = nome_terminal
                row_cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                row_cells[0].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                if row_cells[0].paragraphs[0].runs:
                    _aplicar_estilo_resumo(row_cells[0].paragraphs[0].runs[0], negrito=True)
            
            p_nc = row_cells[1].paragraphs[0]
            aplicar_estilo_paragrafo_compacto(p_nc)
            
            r_id = p_nc.add_run(f"{item['id']}")
            _aplicar_estilo_resumo(r_id, negrito=True)
         
            desc_limpa = _limpar_redundancia_tabela(item['desc'])
            
            # --- AJUSTE AQUI: Trocado '–' (travessão) por ' - ' (hífen) ---
            r_desc = p_nc.add_run(f" - {desc_limpa}") 
            _aplicar_estilo_resumo(r_desc)
            
            row_cells[2].text = item["socicam"]
            aplicar_estilo_paragrafo_compacto(row_cells[2].paragraphs[0])
            if row_cells[2].paragraphs[0].runs: _aplicar_estilo_resumo(row_cells[2].paragraphs[0].runs[0])
            
            row_cells[3].text = item["data"]
            row_cells[3].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            row_cells[3].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_cells[3].paragraphs[0].runs: _aplicar_estilo_resumo(row_cells[3].paragraphs[0].runs[0])
            
            s = str(item["situacao"]).upper().strip()
            row_cells[4].text = s
            row_cells[4].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            row_cells[4].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            
            is_critico = s in ["PENDENTE", "NÃO CONFORME", "NAO CONFORME", "NO PRAZO"]
            if row_cells[4].paragraphs[0].runs:
                _aplicar_estilo_resumo(row_cells[4].paragraphs[0].runs[0], negrito=is_critico)

            current_row_index += 1

        if len(lista_dados) > 1 and primeira_celula_terminal:
            try:
                ultima_celula = tabela.cell(current_row_index - 1, 0)
                primeira_celula_terminal.merge(ultima_celula)
            except: pass

    doc.add_paragraph()