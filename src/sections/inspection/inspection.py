from docx import Document
from docx.shared import Pt, Cm, Inches
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL 
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor 
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml 
from utils import (
    adicionar_titulo_secao, 
    adicionar_paragrafo_justificado, 
    adicionar_titulo_quadro, 
    remover_espacamento_paragrafo,
    extrair_cidade,
    padronizar_processo 
) 
import pandas as pd
import re 

def apply_shading(cell, color_hex="D9D9D9"): 
    """Aplica cor de fundo (shading) a uma célula."""
    shading_elm = parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls('w'), color_hex))
    cell._tc.get_or_add_tcPr().append(shading_elm)


def gerar_secao_fiscalizacao(doc: Document, row, nao_conformidades_df):
    """
    Gera a seção '4. FISCALIZAÇÃO' e a tabela de Não Conformidades (Quadro 1), 
    com informações dinâmicas de equipe e locais.
    """

    adicionar_titulo_secao(doc, "4. FISCALIZAÇÃO")

    doc.add_paragraph() 

    nomes_responsaveis = [n.strip() for n in str(row["Pessoal Responsável"]).split(';') if n.strip()]
    matriculas_responsaveis = [m.strip() for m in str(row["Matrícula do Pessoal Responsável"]).split(';') if m.strip()]

    matriculas_ajustadas = matriculas_responsaveis + [''] * (len(nomes_responsaveis) - len(matriculas_responsaveis))
    responsaveis_info = list(zip(nomes_responsaveis, matriculas_ajustadas))

    terminais_originais = [t.strip() for t in str(row["Terminais"]).split(';') if t.strip()]

    locais_formatados = []

    for terminal_original in terminais_originais:
        cidade_limpa = extrair_cidade(terminal_original)

        if "recife (tip)" in terminal_original.lower():
            locais_formatados.append(f"em Recife (TIP)")
        elif cidade_limpa:
            locais_formatados.append(f"na cidade de {cidade_limpa}")


    par_equipe = doc.add_paragraph()
    remover_espacamento_paragrafo(par_equipe) 
    par_equipe.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    par_equipe.add_run("As ações de fiscalização foram realizadas pela equipe formada pelos Analistas de Regulação ")

    for i, (nome, matricula) in enumerate(responsaveis_info):
        par_equipe.add_run(nome).bold = True

        if matricula:
            par_equipe.add_run(f" (matrícula nº {matricula})") 

        if i < len(responsaveis_info) - 2:
            par_equipe.add_run(", ")
        elif i == len(responsaveis_info) - 2:
            par_equipe.add_run(" e ")

    texto_locais = "; ".join(locais_formatados)
    par_equipe.add_run(f", {texto_locais}.") 

    doc.add_paragraph() 
    par_nc_intro = doc.add_paragraph()
    remover_espacamento_paragrafo(par_nc_intro)
    par_nc_intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    par_nc_intro.add_run("As Não Conformidades constatadas estão relacionadas ao ")
    par_nc_intro.add_run("Programa de Manutenção dos Terminais Rodoviários").bold = True
    par_nc_intro.add_run(", Anexo V do Contrato de Concessão, conforme descritas no ")
 
    run_quadro1 = par_nc_intro.add_run("Quadro 1")
    run_quadro1.bold = True

    par_nc_intro.add_run(", a seguir, com indicação dos respectivos registros fotográficos no ")
    par_nc_intro.add_run("Apêndice 1").bold = True
    par_nc_intro.add_run(".")

    doc.add_paragraph() 

    adicionar_titulo_quadro(
        doc, 
        "Quadro 1 – Não Conformidades por Terminal Rodoviário de Passageiros", 
        negrito=True
    )

    processo_filtragem = str(row["PROCESSO"]).strip()

    if "PROCESSO" not in nao_conformidades_df.columns:
        adicionar_paragrafo_justificado(doc, "⚠️ Coluna 'PROCESSO' não encontrada na planilha de Não-conformidades. Não é possível filtrar os dados.")
        return

    nao_conformidades_df['PROCESSO_LIMPO'] = nao_conformidades_df['PROCESSO'].astype(str).str.strip()

    nc_fisc = nao_conformidades_df[
        nao_conformidades_df["PROCESSO_LIMPO"] == processo_filtragem
    ].copy()

    if nc_fisc.empty:

        adicionar_paragrafo_justificado(doc, f"Nenhuma não conformidade registrada para o processo buscado: **{processo_filtragem}**. Verifique a planilha.")
        return
        
    total_nc = len(nc_fisc)

    colunas_necessarias = ["TERMINAL RODOVIÁRIO", "ID", "DESCRIÇÃO", "REGISTROS FOTOGRÁFICOS", "FUNDAMENTO DA INFRAÇÃO", "DETERMINAÇÃO"]

    for col in colunas_necessarias:
       
        if col not in nc_fisc.columns:
            print(f"⚠️ Coluna '{col}' não encontrada na planilha de Não-conformidades.")

    tabela = doc.add_table(rows=1, cols=6) 
    tabela.style = "Table Grid"
    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER 

    tabela.columns[0].width = Cm(1.2) # TRP
    tabela.columns[1].width = Cm(2.5) # IDENTIFICAÇÃO
    tabela.columns[2].width = Cm(4.5) # DESCRIÇÃO
    tabela.columns[3].width = Cm(2.5) # REGISTRO FOTOGRÁFICO
    tabela.columns[4].width = Cm(3.8) # FUNDAMENTO DA INFRAÇÃO
    tabela.columns[5].width = Cm(2.5) # DETERMINAÇÃO

    # Cabeçalhos
    cabecalhos_nomes = ["TRP", "IDENTIFICAÇÃO", "DESCRIÇÃO", "REGISTRO FOTOGRÁFICO", "FUNDAMENTO DA INFRAÇÃO (ANEXO V CONTRATO DE CONCESSÃO)", "DETERMINAÇÃO"]

    cabecalho_cells = tabela.rows[0].cells

    for i, nome in enumerate(cabecalhos_nomes):
        cell = cabecalho_cells[i]
      
        apply_shading(cell) 

        cell.text = nome
        
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER 

        for par in cell.paragraphs:
            remover_espacamento_paragrafo(par)
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in par.runs:
                run.bold = True
                run.font.size = Pt(9) 

    nc_fisc = nc_fisc.sort_values(by=["TERMINAL RODOVIÁRIO", "ID"])

    # Dicionário para rastrear a primeira linha de cada TRP
    linhas_terminais = {} 

    # Criar todas as linhas primeiro
    for idx, linha in nc_fisc.iterrows():
        row_cells = tabela.add_row().cells

        # Colunas com dados
        data_map = {
            1: linha.get("ID", ""),
            2: linha.get("DESCRIÇÃO", ""),
            3: linha.get("REGISTROS FOTOGRÁFICOS", ""),
            4: linha.get("FUNDAMENTO DA INFRAÇÃO", ""),
            5: linha.get("DETERMINAÇÃO", ""),
         }

        trp_sigla = linha.get("TERMINAL RODOVIÁRIO", "")

        for i, (col_index, text) in enumerate(data_map.items()):
            par = row_cells[col_index].paragraphs[0]
            remover_espacamento_paragrafo(par)
            par.text = text
           
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if col_index == 2 else WD_ALIGN_PARAGRAPH.LEFT
            for run in par.runs:
                run.font.size = Pt(10)

            row_cells[col_index].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

        if trp_sigla not in linhas_terminais:
            linhas_terminais[trp_sigla] = tabela.rows[-1].cells[0] 

        par = row_cells[0].paragraphs[0]
        remover_espacamento_paragrafo(par)
        if trp_sigla == nc_fisc.iloc[nc_fisc.index.get_loc(idx) - 1].get("TERMINAL RODOVIÁRIO") and nc_fisc.index.get_loc(idx) != 0:
            par.text = "" 
        else:
            par.text = trp_sigla 
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in par.runs:
                run.font.size = Pt(10)

    start_row = 1 

    for trp, group in nc_fisc.groupby("TERMINAL RODOVIÁRIO"):

        count_rows = len(group)
        end_row = start_row + count_rows - 1

        if count_rows > 1:
            primeira_celula = tabela.cell(start_row, 0)
            ultima_celula = tabela.cell(end_row, 0)

            merged_cell = primeira_celula.merge(ultima_celula)

            merged_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for par in merged_cell.paragraphs:
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER

        start_row = end_row + 1
  
    total_row_cells = tabela.add_row().cells
  
    merged_total_cell = total_row_cells[0].merge(total_row_cells[4])
    
    apply_shading(merged_total_cell)
    
    par_total = merged_total_cell.paragraphs[0]
    remover_espacamento_paragrafo(par_total)
    par_total.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_total = par_total.add_run("TOTAL")
    run_total.bold = True
    run_total.font.size = Pt(10)
    
    cell_count = total_row_cells[5]
    
    apply_shading(cell_count)
    cell_count.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    
    par_count = cell_count.paragraphs[0]
    remover_espacamento_paragrafo(par_count)
    par_count.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_count = par_count.add_run(str(total_nc))
    run_count.bold = True
    run_count.font.size = Pt(10)

    doc.add_paragraph()