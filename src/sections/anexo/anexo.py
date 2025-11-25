from typing import List, Optional, Dict
from itertools import zip_longest
import os
import re
import pandas as pd
from docx.document import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from utils import (
    adicionar_titulo_secao,
    adicionar_duas_imagens_lado_a_lado,
    aplicar_estilo_corpo,
    _adicionar_contexto_nc,
    encontrar_dados_na_base
)

def _safe_split_nc_separator(content: str) -> List[str]:
    """
    Divide apenas por ';' ou quebra de linha.
    Usado para separar os blocos de NCs e Legendas principais.
    """
    if not isinstance(content, str) or content.lower() == "nan" or not content.strip():
        return []
    return [d.strip() for d in re.split(r'[;\n]+', content) if d.strip()]

def _adaptar_lista_anexo(lista: List[str], tamanho_alvo: int, default_val: str) -> List[str]:
    """Garante que a lista tenha o tamanho alvo, preenchendo ou cortando."""
    if len(lista) == tamanho_alvo: return lista
    if len(lista) > tamanho_alvo: return lista[:tamanho_alvo]
    nova = list(lista)
    nova.extend([default_val] * (tamanho_alvo - len(lista)))
    return nova

def _extrair_cidades_desde_nc(nc_fisc_df: pd.DataFrame) -> List[str]:
    cidades = []
    seen = set()
    if "Terminal" not in nc_fisc_df.columns: return []
    for raw in nc_fisc_df["Terminal"].fillna("").astype(str).tolist():
        s = raw.strip()
        if not s: continue
        s = re.sub(r"(?i)^\s*terminal\s+de\s+", "", s)
        s = re.sub(r"\s*\(.*?\)\s*$", "", s).strip()
        if s and s not in seen:
            cidades.append(s)
            seen.add(s)
    return cidades

def _formatar_periodos(periodo_raw: Optional[str]) -> Optional[str]:
    if not isinstance(periodo_raw, str) or not periodo_raw.strip() or periodo_raw.lower() == "nan":
        return None
    partes = [p.strip() for p in periodo_raw.split(";") if p.strip()]
    if not partes: return None
    if len(partes) == 1: return partes[0]
    return " e ".join(partes)

def _buscar_arquivos_flexivel(id_excel: str, lista_arquivos: List[str]) -> List[str]:
    """
    Busca arquivos (Lógica Simples Original).
    """
    id_limpo = id_excel.strip().upper()
    
    matches = [f for f in lista_arquivos if f.upper().startswith(id_limpo)]
    if matches: return sorted(matches)
    
    if "." in id_limpo:
        id_base = id_limpo.rsplit(".", 1)[0]
        matches = [f for f in lista_arquivos if f.upper().startswith(id_base)]
        if matches: return sorted(matches)

    id_norm = re.sub(r"[_\-\s\.]", "", id_limpo)
    matches_norm = []
    for f in lista_arquivos:
        f_nome = os.path.splitext(f)[0]
        f_norm = re.sub(r"[_\-\s\.]", "", f_nome).upper()
        if f_norm.startswith(id_norm):
            matches_norm.append(f)
    if matches_norm: return sorted(matches_norm)

    if len(id_limpo) > 5:
        matches_contains = [f for f in lista_arquivos if id_limpo in f.upper()]
        if matches_contains: return sorted(matches_contains)

    return []

def _limpar_redundancia_anexo(texto: str) -> str:
    """Remove 'TIP 01 - ' do início do título da foto."""
    if not texto: return ""
    padrao = r'^(?:[A-Z]{3}[\s\-]*)?\d+(?:[._]\d+)?\s*[-:–)]*\s*'
    texto_limpo = re.sub(padrao, '', str(texto).strip())
    
    if not texto_limpo and texto: return str(texto) 
    
    if texto_limpo and texto_limpo[0].islower():
        return texto_limpo[0].upper() + texto_limpo[1:]
    
    return texto_limpo

def gerar_secao_anexo_fotos(
    doc: Document,
    row: pd.Series,
    nao_conformidades_df: pd.DataFrame,
    fotos_dir: str,
    processo_info: dict,
    df_base_nc: pd.DataFrame,
    ano_user: str,
    proc_user: str,
    monit_user: str
) -> None:
    doc.add_page_break()
    periodo_raw = processo_info.get("Periodo de Vistoria da ARPE", None)
    periodo_formatado = _formatar_periodos(periodo_raw)
    processo_ctr_num = processo_info.get("Processo CTR Nº", "XX/XXXX")
    id_fisc = row["ID da Fiscalização"]
    nc_fisc = nao_conformidades_df[nao_conformidades_df["ID da Fiscalização"] == id_fisc].copy()

    cidades = _extrair_cidades_desde_nc(nc_fisc)
    cidades_str = ", ".join(cidades) if cidades else "(NOMES DAS CIDADES)"
    titulo_base = "ANEXO - MEMORIAL FOTOGRÁFICO - VISTORIAS REALIZADAS"
    titulo_completo = f"{titulo_base} EM {periodo_formatado}" if periodo_formatado else f"{titulo_base} EM DATA INDEFINIDA"

    adicionar_titulo_secao(doc, titulo_completo, aplicar_sombra=True)

    if nc_fisc.empty:
        doc.add_paragraph("Nenhuma não conformidade monitorada disponível.")
        return

    paragrafo_anexo = doc.add_paragraph()
    paragrafo_anexo.paragraph_format.space_after = Pt(12)
    texto_inicial = (
        "Apresenta-se, a seguir, evidências fotográficas das Não Conformidades pendentes do "
        f"Relatório de Fiscalização Técnico-Operacional Arpe/CTR nº {processo_ctr_num} para os Terminais Rodoviários "
        f"de Passageiros concedidos à SOCICAM nas cidades de {cidades_str}."
    )
    run_paragrafo = paragrafo_anexo.add_run(texto_inicial)
    aplicar_estilo_corpo(run_paragrafo, negrito=False)
    paragrafo_anexo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY_LOW

    if not os.path.exists(fotos_dir):
        doc.add_paragraph(f"⚠️ ERRO: Pasta de fotos não encontrada: {fotos_dir}")
        return

    try:
        arquivos_na_pasta = [f for f in os.listdir(fotos_dir) if f.lower().endswith(('.jpg', '.jpeg', '.png'))]
    except Exception as e:
        doc.add_paragraph(f"⚠️ ERRO ao ler pasta de fotos: {e}")
        return

    terminal_anterior = None

    for terminal, grupo_terminal in nc_fisc.groupby("Terminal"):
        
        # --- CORREÇÃO: Conjunto para evitar IDs duplicados no mesmo terminal ---
        ids_processados_no_terminal = set() 

        for _, linha in grupo_terminal.iterrows():
            
            raw_key = str(linha.get("Legenda da Foto", "")).strip()
            if not raw_key or raw_key.lower() == "nan":
                 raw_key = str(linha.get("Constatação", "")).strip()
                 if not raw_key or raw_key.lower() == "nan":
                    raw_key = str(linha.get("Não Conformidade", "")).strip()

            textos_busca = _safe_split_nc_separator(raw_key)
    
            raw_legenda = str(linha.get("Legenda da Foto", "")).strip()
            if raw_legenda.lower() == "nan": raw_legenda = ""
            legendas_em_bloco = _safe_split_nc_separator(raw_legenda)
       
            max_len = len(textos_busca)
            legendas_em_bloco = _adaptar_lista_anexo(legendas_em_bloco, max_len, "")

            for i in range(max_len):
                texto_id = textos_busca[i]
                bloco_legenda = legendas_em_bloco[i] 
                
                if not texto_id: continue

                # --- LIMPEZA DE BUSCA (IGUAL AO NONCONFORMITY.PY) ---
                # Pega só o texto antes dos dois pontos para buscar o ID correto
                texto_para_busca = texto_id.split(':')[0].strip()

                # Busca ID na Base
                dados_base = encontrar_dados_na_base(
                    texto_para_busca, 
                    df_base_nc, 
                    ano_user, 
                    proc_user, 
                    monit_user,
                    terminal_user=str(terminal)
                )
                
                id_encontrado = dados_base["id"]
                descricao_oficial = dados_base["descricao"]

                if id_encontrado in ["ID_NAO_ENCONTRADO", "ID_ERRO"]:
                    continue

                # --- CORREÇÃO: Verifica se já processamos esse ID neste terminal ---
                if id_encontrado in ids_processados_no_terminal:
                    continue
                
                # Se não processou, adiciona na lista para não repetir
                ids_processados_no_terminal.add(id_encontrado)

                # Busca fotos
                fotos_do_item = _buscar_arquivos_flexivel(id_encontrado, arquivos_na_pasta)

                if not fotos_do_item: continue

                # Ordena
                fotos_do_item.sort()

                if terminal != terminal_anterior:
                    t_limpo = re.sub(r"(?i)^\s*terminal\s+de\s+", "", str(terminal))
                    t_limpo = re.sub(r"\s*\(.*?\)\s*$", "", t_limpo).strip()
                    doc.add_paragraph().paragraph_format.space_after = Pt(12)
                    adicionar_titulo_secao(doc, f"TERMINAL DE {t_limpo.upper()}", nivel_heading=2)
                    terminal_anterior = terminal

                # Limpeza do Título da NC
                desc_curta = descricao_oficial.split('\n')[0]
                desc_limpa = _limpar_redundancia_anexo(desc_curta)
                _adicionar_contexto_nc(doc, id_encontrado, desc_limpa)

                if ":" in bloco_legenda:
                    sub_legendas = [x.strip() for x in bloco_legenda.split(":")]
                else:
                    sub_legendas = [bloco_legenda] * len(fotos_do_item)

                sub_legendas = _adaptar_lista_anexo(sub_legendas, len(fotos_do_item), "")
                
                fotos_em_pares = list(zip_longest(*[iter(fotos_do_item)] * 2, fillvalue=None))

                count_legenda = 0
                for foto1, foto2 in fotos_em_pares:
                    l1 = sub_legendas[count_legenda] if count_legenda < len(sub_legendas) else ""
                    count_legenda += 1
                    l2 = ""
                    if foto2:
                        l2 = sub_legendas[count_legenda] if count_legenda < len(sub_legendas) else ""
                        count_legenda += 1
                    
                    adicionar_duas_imagens_lado_a_lado(
                        doc, fotos_dir, foto1, l1, foto2, l2, contexto_nc_tupla=None
                    )
    
    doc.add_page_break()