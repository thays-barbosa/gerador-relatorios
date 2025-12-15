from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from openpyxl import load_workbook
import os
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
import io
from docx import Document
import pandas as pd
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.table import WD_TABLE_ALIGNMENT
from datetime import datetime
import re

# --- 1. NOVAS FUNÇÕES PARA O SUMÁRIO (INDISPENSÁVEIS) ---

def adicionar_titulo_secao(doc, texto, nivel=1):
    """
    MODIFICADA: Agora usa o 'level' nativo do Word para o Sumário Automático funcionar.
    """
    titulo = doc.add_heading(texto, level=nivel)
    for run in titulo.runs:
        run.font.name = 'Arial'
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.bold = True
    titulo.alignment = WD_ALIGN_PARAGRAPH.LEFT
    titulo.paragraph_format.space_before = Pt(12)
    titulo.paragraph_format.space_after = Pt(6)

def forcar_atualizacao_campos(doc):
    """Configura o Word para atualizar o Sumário (TOC) ao abrir."""
    element = doc.settings.element
    update_fields = OxmlElement('w:updateFields')
    update_fields.set(qn('w:val'), 'true')
    element.append(update_fields)

# --- 2. FUNÇÕES DE FORMATAÇÃO (SUAS ORIGINAIS INTEGRALMENTE) ---

def remover_espacamento_paragrafo(paragrafo):
    paragrafo_format = paragrafo.paragraph_format
    paragrafo_format.space_before = Pt(0)
    paragrafo_format.space_after = Pt(0)

def adicionar_quebra_linha_controlada(doc, altura_pt=18):
    paragrafo = doc.add_paragraph()
    remover_espacamento_paragrafo(paragrafo)
    run = paragrafo.add_run("")
    run.font.size = Pt(altura_pt)

def adicionar_paragrafo_justificado(doc, texto, tamanho_fonte=12):
    paragrafo = doc.add_paragraph(texto)
    paragrafo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for run in paragrafo.runs:
        run.font.name = 'Arial'
        run.font.size = Pt(tamanho_fonte)

def adicionar_texto_centralizado(doc, texto, tamanho_fonte=12, negrito=True):
    paragraph = doc.add_paragraph()
    remover_espacamento_paragrafo(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(texto)
    run.bold = negrito
    run.font.name = 'Arial'
    run.font.size = Pt(tamanho_fonte)

def adicionar_titulo_quadro(doc, texto, negrito=False, tamanho=11):
    paragrafo = doc.add_paragraph()
    paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    remover_espacamento_paragrafo(paragrafo)
    run = paragrafo.add_run(texto)
    run.bold = negrito
    run.font.name = 'Arial'
    run.font.size = Pt(tamanho)

def aplicar_estilo_texto(run, tamanho=12, negrito=False, fonte="Arial", cor_rgb=(0, 0, 0)):
    run.font.name = fonte
    run._element.rPr.rFonts.set(qn("w:eastAsia"), fonte)
    run.font.size = Pt(tamanho)
    run.bold = negrito
    run.font.color.rgb = RGBColor(*cor_rgb)

def aplicar_borda_paragrafo(paragraph):
    p = paragraph._element
    pPr = p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    for border_name in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{border_name}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "2")
        border.set(qn("w:color"), "000000")
        borders.append(border)
    pPr.append(borders)

def adicionar_legenda_formatada(doc, texto):
    par = doc.add_paragraph()
    remover_espacamento_paragrafo(par)
    run = par.add_run(texto)
    aplicar_estilo_texto(run, tamanho=10, fonte="Arial", cor_rgb=(90, 90, 90))
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    aplicar_borda_paragrafo(par)

# --- 3. IMAGENS E APÊNDICE FOTOGRÁFICO (SUAS ORIGINAIS) ---

def processar_imagem_para_relatorio(caminho_imagem, largura_max=1024, qualidade=80):
    img = Image.open(caminho_imagem)
    if img.mode != "RGB":
        img = img.convert("RGB")
    if img.width > largura_max:
        proporcao = largura_max / float(img.width)
        altura_nova = int(float(img.height) * proporcao)
        img = img.resize((largura_max, altura_nova), Image.Resampling.LANCZOS)
    buffer = io.BytesIO()
    img.save(buffer, format="JPEG", quality=qualidade, optimize=True)
    buffer.seek(0)
    return buffer

def adicionar_imagem_na_celula(celula, caminho_imagem, largura_max_cm=7.5):
    paragrafo_img = celula.add_paragraph()
    remover_espacamento_paragrafo(paragrafo_img)
    paragrafo_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    imagem_buffer = processar_imagem_para_relatorio(caminho_imagem)
    run = paragrafo_img.add_run()
    run.add_picture(imagem_buffer, width=Cm(largura_max_cm))
    celula.vertical_alignment = WD_ALIGN_VERTICAL.TOP

def adicionar_apendice_fotos(doc, caminho_base_fotos, id_fiscalizacao, caminho_planilha_legendas):
    aba_nc = "Não-conformidades " 
    try:
        df_legendas = pd.read_excel(caminho_planilha_legendas, sheet_name=aba_nc)
    except ValueError:
        try:
            df_legendas = pd.read_excel(caminho_planilha_legendas, sheet_name=aba_nc.strip())
        except Exception:
             adicionar_paragrafo_justificado(doc, f"ERRO: Aba de legendas não encontrada.")
             return

    df_legendas.columns = df_legendas.columns.str.strip() 
    id_fisc_limpo = str(id_fiscalizacao).strip()
    linha_legenda = df_legendas[df_legendas['ID da Fiscalização'].astype(str).str.strip() == id_fisc_limpo]
     
    if linha_legenda.empty:
        legendas = [] 
    else:
        texto_legendas = str(linha_legenda['Legenda da Foto'].iloc[0]) 
        legendas = [l.strip() for l in texto_legendas.split(';') if l.strip()]

    try:
        arquivos_fotos = sorted([os.path.join(caminho_base_fotos, f) for f in os.listdir(caminho_base_fotos) if f.lower().endswith(('.png', '.jpg', '.jpeg'))])
    except: return

    num_fotos = len(arquivos_fotos)
    num_linhas_tabela = ((num_fotos + 1) // 2) * 2 
    tabela = doc.add_table(rows=num_linhas_tabela, cols=2)
    tabela.style = 'Table Grid'

    for i in range(num_linhas_tabela // 2):
        for j in range(2):
            foto_idx = i * 2 + j
            if foto_idx < num_fotos:
                adicionar_imagem_na_celula(tabela.cell(i*2, j), arquivos_fotos[foto_idx])
                celula_legenda = tabela.cell(i*2+1, j)
                p_leg = celula_legenda.paragraphs[0]
                remover_espacamento_paragrafo(p_leg)
                p_leg.alignment = WD_ALIGN_PARAGRAPH.CENTER
                txt = legendas[foto_idx] if foto_idx < len(legendas) else f"Foto {foto_idx+1}"
                aplicar_estilo_texto(p_leg.add_run(txt), tamanho=10, cor_rgb=(90, 90, 90))

# --- 4. TABELAS (SUAS ORIGINAIS INTEGRALMENTE) ---

def aplicar_fundo_cinza(celula):
    shading_elm = OxmlElement('w:shd')
    shading_elm.set(qn('w:val'), 'clear')
    shading_elm.set(qn('w:fill'), 'DDDDDD') 
    celula._element.tcPr.append(shading_elm)

def adicionar_cabecalho_tabela(doc, texto, tamanho_fonte=12):
    paragrafo = doc.add_paragraph()
    remover_espacamento_paragrafo(paragrafo)
    run = paragrafo.add_run(texto.upper())
    run.bold = True
    paragrafo.alignment = WD_ALIGN_PARAGRAPH.LEFT 
    adicionar_quebra_linha_controlada(doc, altura_pt=12)

def adicionar_tabela_informacoes(doc, dados_tabela):
    adicionar_cabecalho_tabela(doc, "INFORMAÇÕES GERAIS")
    tabela = doc.add_table(rows=len(dados_tabela), cols=2)
    tabela.autofit = False
    tabela.columns[0].width = Cm(6.0) 
    tabela.columns[1].width = Cm(11.0)
    tabela.style = 'Table Grid'
    titulos_cinza = ["3.1 DO TITULAR", "3.2 DO REGULADO", "3.3 DO REGULADOR"]
    for i, (rotulo, valor) in enumerate(dados_tabela):
        cel_r, cel_v = tabela.rows[i].cells[0], tabela.rows[i].cells[1]
        if rotulo in titulos_cinza:
            aplicar_fundo_cinza(cel_r); aplicar_fundo_cinza(cel_v)
            p = cel_r.merge(cel_v).paragraphs[0]
            remover_espacamento_paragrafo(p); p.add_run(rotulo).bold = True
            continue
        pr = cel_r.paragraphs[0]; remover_espacamento_paragrafo(pr); pr.add_run(rotulo).bold = True
        pv = cel_v.paragraphs[0]; remover_espacamento_paragrafo(pv); run_v = pv.add_run(str(valor))
        if rotulo in ["Responsável:", "Diretor Presidente:"]: run_v.bold = True

def adicionar_tabela_abreviaturas(doc, df_abreviaturas):
    tabela = doc.add_table(rows=len(df_abreviaturas) + 1, cols=2)
    tabela.style = 'Table Grid'
    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabela.columns[0].width = Cm(2.5); tabela.columns[1].width = Cm(14.5)
    for i, txt in enumerate(["SIGLA", "DEFINIÇÃO"]):
        p = tabela.rows[0].cells[i].paragraphs[0]
        remover_espacamento_paragrafo(p); p.add_run(txt).bold = True; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for i, row in df_abreviaturas.iterrows():
        c0, c1 = tabela.rows[i+1].cells[0], tabela.rows[i+1].cells[1]
        p0 = c0.paragraphs[0]; remover_espacamento_paragrafo(p0); p0.add_run(str(row['Sigla'])); p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1 = c1.paragraphs[0]; remover_espacamento_paragrafo(p1); p1.add_run(str(row['Definição']))

# --- 5. LOGICA EXCEL E DATA (SUAS ORIGINAIS INTEGRALMENTE) ---

def ajustar_largura_colunas(caminho_planilha):
    wb = load_workbook(caminho_planilha)
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for coluna in ws.columns:
            max_length = 0
            for celula in coluna:
                if celula.value: max_length = max(max_length, len(str(celula.value)))
            ws.column_dimensions[coluna[0].column_letter].width = max(max_length + 2, 10)
    wb.save(caminho_planilha)

def arquivo_em_uso(caminho):
    try: os.rename(caminho, caminho); return False
    except PermissionError: return True

def parear_responsaveis_e_matriculas(nomes_str, matriculas_str, cargo_fixo="Analista de Regulação, matrícula nº"):
    if not nomes_str or not matriculas_str: return []
    nomes = [n.strip() for n in str(nomes_str).split(';') if n.strip()]
    matriculas = [m.strip() for m in str(matriculas_str).split(';') if m.strip()]
    return [{'nome': n, 'info_completa': f"{cargo_fixo} {m}"} for n, m in zip(nomes, matriculas)]

def formatar_data_capa(data_str):
    if not data_str: return ""
    try:
        meses = {'01':'Janeiro','02':'Fevereiro','03':'Março','04':'Abril','05':'Maio','06':'Junho','07':'Julho','08':'Agosto','09':'Setembro','10':'Outubro','11':'Novembro','12':'Dezembro'}
        p = str(data_str).split('/')
        return f"{meses.get(p[1], p[1])}, {p[2]}"
    except: return str(data_str)

def parear_e_formatar_assinaturas(nomes_str, matriculas_str, cargo_fixo="Analista de Regulação"):
    nomes = [n.strip() for n in str(nomes_str).split(';') if n.strip()]
    matriculas = [m.strip() for m in str(matriculas_str).split(';') if m.strip()]
    return [(n, cargo_fixo, f"Matrícula nº {matriculas[i]}" if i < len(matriculas) else "") for i, n in enumerate(nomes)]

def extrair_cidade(terminal_str):
    if not isinstance(terminal_str, str): return ""
    c = re.sub(r'terminal d[eo]\s*', '', terminal_str, flags=re.IGNORECASE)
    c = re.sub(r'\(TIP\)', ' (TIP)', c, flags=re.IGNORECASE).strip() 
    c = re.sub(r'\s*\(.*\)', '', c).strip()
    return c.capitalize()

def padronizar_processo(id_fisc):
    if not isinstance(id_fisc, str): return ""
    partes = id_fisc.upper().replace('-', ' ').split()
    return f"{partes[0]} {partes[1]}/{partes[2]}" if len(partes) == 3 else id_fisc

# --- 6. ASSINATURAS (SUA LÓGICA VERTICAL ORIGINAL) ---

def adicionar_assinaturas_formatadas(doc, analistas_fixos, coordenador_nome_fixo, cidade_relatorio="Recife"):
    data_par = doc.add_paragraph()
    remover_espacamento_paragrafo(data_par)
    data_par.paragraph_format.space_before = Pt(36)
    data_par.alignment = WD_ALIGN_PARAGRAPH.LEFT 
    data_par.add_run(f"{cidade_relatorio}, data da assinatura eletrônica.").font.size = Pt(12)
    doc.add_paragraph().paragraph_format.space_after = Pt(72)

    def _bloco(doc, nome, cargo, matricula):
        p = doc.add_paragraph(); remover_espacamento_paragrafo(p); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(nome).bold = True; p.runs[0].font.size = Pt(11)
        doc.add_paragraph(cargo).alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph(matricula).alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph().paragraph_format.space_after = Pt(36)

    for analista in analistas_fixos:
        _bloco(doc, analista[0], analista[1], analista[2])

    p_ciente = doc.add_paragraph()
    remover_espacamento_paragrafo(p_ciente); p_ciente.paragraph_format.space_before = Pt(72); p_ciente.alignment = WD_ALIGN_PARAGRAPH.LEFT 
    p_ciente.add_run("Ciente e de acordo.").font.size = Pt(11)
    doc.add_paragraph().paragraph_format.space_after = Pt(72)
    
    p_n = doc.add_paragraph(); remover_espacamento_paragrafo(p_n); p_n.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_n.add_run(coordenador_nome_fixo).bold = True
    doc.add_paragraph("Coordenadora de Transportes e Rodovias").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Matrícula nº 209640/01").alignment = WD_ALIGN_PARAGRAPH.CENTER